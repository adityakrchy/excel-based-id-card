"""Excel and CSV data source importer supporting .xlsx, .xls, and .csv."""

import logging
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any, Callable, Dict, List, Optional, Tuple
import pandas as pd

from id_card_generator.core.exceptions import DataSourceError, SheetNotFoundError
from id_card_generator.core.models import SourceRecord
from id_card_generator.data.data_normalizer import DataNormalizer
from id_card_generator.data.hindi_converter import HindiConverter

logger = logging.getLogger(__name__)

# Progress callback signature: (completed, total, message)
LoadProgressCallback = Optional[Callable[[int, int, str], None]]
LoadCancelCheck = Optional[Callable[[], bool]]


@dataclass
class ExcelSheetMatrix:
    """Raw worksheet content captured in a single pass (Stage A of the load pipeline).

    Attributes:
        columns: Converted column names (Hindi-aware), as produced by ``import_data``.
        rows: Data rows below the header; every cell keeps its raw value
            (``None`` marks the empty cells that ``import_data`` sees as NaN).
        row_index_offset: The ``record.row_index`` of the first data row
            (``header_row + 1``), matching the historical ``import_data`` convention.
        pin_col: Detected PIN column name (may be ``None`` when the sheet is empty).
        row_fonts: Sparse map ``{row_index: {column_index: legacy_font_name}}``
            containing only cells whose font is a legacy Devanagari font, so the
            ``auto`` Hindi mode can pick them up. Every other cell resolves to
            ``None`` which reproduces the non-legacy code path of ``import_data``.
    """

    columns: List[str] = field(default_factory=list)
    rows: List[List[Any]] = field(default_factory=list)
    row_index_offset: int = 0
    pin_col: Optional[str] = None
    row_fonts: Dict[int, Dict[int, str]] = field(default_factory=dict)

    @property
    def total_rows(self) -> int:
        return len(self.rows)


def _progress_shim(progress_cb: LoadProgressCallback, done: int, total: int, msg: str = "") -> None:
    """Guard against progress callbacks throwing inside the importer."""
    if progress_cb:
        try:
            progress_cb(done, total, msg)
        except Exception:  # pragma: no cover - defensive
            logger.debug("Progress callback failed", exc_info=True)


def _cancelled(cancel_check: LoadCancelCheck) -> bool:
    """Compatibility guard: a raising cancellation check is treated as 'no'."""
    if cancel_check:
        try:
            return bool(cancel_check())
        except Exception:  # pragma: no cover - defensive
            logger.debug("Cancellation callback failed", exc_info=True)
    return False


def convert_chunk_job(job: Dict[str, Any]) -> Tuple[int, int, List[SourceRecord]]:
    """Convert one chunk of raw worksheet rows into ``SourceRecord`` objects.

    Module-level function so it can be pickled for the process-pool backend
    (feature flag ``ID_STREAM_PROCESS_POOL=1``).

    Args:
        job: Dictionary with the keys ``start``, ``rows``, ``row_fonts``,
            ``columns``, ``pin_col``, ``row_index_offset`` and ``hindi_mode``.

    Returns:
        Tuple of ``(start, record_count, records)``.
    """
    records = ExcelImporter.build_chunk_records(
        row_values=job["rows"],
        row_fonts=job.get("row_fonts") or {},
        columns=job["columns"],
        pin_col=job.get("pin_col"),
        first_record_row=job["row_index_offset"],
        hindi_font_mode=job.get("hindi_mode", "auto"),
        start=job.get("start", 0),
    )
    start = int(job["start"])
    return start, len(records), records


class LoadCancelled(Exception):
    """Internal signal: loading was cancelled during the file pass."""


class ExcelImporter:
    """Imports spreadsheets and CSV files with sheet selection, header configuration, and Hindi font processing."""

    @staticmethod
    def get_sheet_names(file_path: Path) -> List[str]:
        """Return list of sheet names in an Excel file. Returns ['Default'] for CSV."""
        if not file_path.exists():
            raise DataSourceError(f"File not found: {file_path}")

        suffix = file_path.suffix.lower()
        if suffix == ".csv":
            return ["Default"]

        try:
            engine = "xlrd" if suffix == ".xls" else "openpyxl"
            with pd.ExcelFile(file_path, engine=engine) as excel_file:
                return list(excel_file.sheet_names)
        except Exception as e:
            logger.error("Failed to inspect sheets in %s: %s", file_path, e)
            raise DataSourceError(f"Could not read spreadsheet sheets: {e}") from e

    @staticmethod
    def import_data(
        file_path: Path,
        sheet_name: Optional[str] = None,
        header_row: int = 0,
        pin_column_hint: Optional[str] = None,
        hindi_font_mode: str = "auto",
    ) -> Tuple[List[str], List[SourceRecord]]:
        """Import records from file.
        
        Args:
            file_path: Path to .xlsx, .xls, or .csv
            sheet_name: Specific sheet name (for Excel workbooks)
            header_row: 0-indexed row number containing column headers
            pin_column_hint: Specific column name to treat as PIN Number (optional)
            hindi_font_mode: 'auto' (detect Kruti Dev / DevLys / Unicode),
                             'kruti_to_unicode' (force conversion), or
                             'none' (preserve raw strings).

        Returns:
            Tuple of (columns_list, source_records_list)
        """
        if not file_path.exists():
            raise DataSourceError(f"File does not exist: {file_path}")

        suffix = file_path.suffix.lower()

        # Extract font metadata if available (for .xlsx)
        header_fonts: Dict[int, str] = {}
        cell_fonts: Dict[Tuple[int, int], str] = {}
        if suffix == ".xlsx":
            try:
                import openpyxl
                wb = openpyxl.load_workbook(file_path, data_only=True)
                ws = wb[sheet_name] if sheet_name and sheet_name in wb.sheetnames else wb.active
                excel_header_row = header_row + 1
                for c_idx in range(1, ws.max_column + 1):
                    cell = ws.cell(excel_header_row, c_idx)
                    if cell.font and cell.font.name:
                        header_fonts[c_idx - 1] = cell.font.name

                for r_idx in range(excel_header_row + 1, ws.max_row + 1):
                    df_r = r_idx - (excel_header_row + 1)
                    for c_idx in range(1, ws.max_column + 1):
                        cell = ws.cell(r_idx, c_idx)
                        if cell.font and cell.font.name:
                            cell_fonts[(df_r, c_idx - 1)] = cell.font.name
            except Exception as e:
                logger.debug("Could not read openpyxl font metadata from %s: %s", file_path, e)

        try:
            # dtype=object preserves original representations (like leading zeros)
            if suffix == ".csv":
                df = pd.read_csv(file_path, header=header_row, dtype=object)
            elif suffix in (".xls", ".xlsx"):
                engine = "xlrd" if suffix == ".xls" else "openpyxl"
                sheet = sheet_name or 0
                with pd.ExcelFile(file_path, engine=engine) as excel_file:
                    df = pd.read_excel(excel_file, sheet_name=sheet, header=header_row, dtype=object)
            else:
                raise DataSourceError(f"Unsupported file format: {suffix}")
        except SheetNotFoundError:
            raise
        except Exception as e:
            logger.error("Failed to import spreadsheet %s: %s", file_path, e)
            raise DataSourceError(f"Failed to read file: {e}") from e

        # Clean and convert column names
        converted_columns: List[str] = []
        for c_idx, raw_col in enumerate(df.columns):
            c_str = str(raw_col).strip()
            # If the column was auto-generated unnamed index, keep as is
            if c_str.startswith("Unnamed:"):
                converted_columns.append(c_str)
            else:
                font_name = header_fonts.get(c_idx)
                conv_col = HindiConverter.convert_value(c_str, font_name=font_name, mode=hindi_font_mode)
                converted_columns.append(conv_col.strip())

        df.columns = converted_columns

        # Identify candidate PIN column
        pin_col = ExcelImporter._detect_pin_column(df.columns.tolist(), pin_column_hint)

        records: List[SourceRecord] = []
        for idx, row in df.iterrows():
            row_dict: Dict[str, Any] = {}
            for c_idx, col in enumerate(df.columns):
                val = row[col]
                # Convert NaN to empty string
                if pd.isna(val):
                    row_dict[col] = ""
                else:
                    font_name = cell_fonts.get((int(idx), c_idx))
                    conv_val = HindiConverter.convert_value(val, font_name=font_name, mode=hindi_font_mode)
                    row_dict[col] = conv_val

            raw_pin_val = row_dict.get(pin_col, "") if pin_col else ""
            raw_pin, normalized_pin = DataNormalizer.normalize_pin(raw_pin_val)

            rec = SourceRecord(
                row_index=int(idx) + header_row + 1,
                raw_pin=raw_pin,
                pin=normalized_pin,
                data=row_dict,
            )
            records.append(rec)

        logger.info(
            "Imported %d records with %d columns from %s (sheet: %s, hindi_mode: %s)",
            len(records),
            len(df.columns),
            file_path.name,
            sheet_name or "default",
            hindi_font_mode,
        )
        return list(df.columns), records

    # ------------------------------------------------------------------
    # Stage A: single-pass raw matrix capture (xlsx single read; CSV/XLS cheap)
    # ------------------------------------------------------------------
    @staticmethod
    def load_matrix(
        file_path: Path,
        sheet_name: Optional[str] = None,
        header_row: int = 0,
        hindi_font_mode: str = "auto",
        pin_column_hint: Optional[str] = None,
        progress_callback: LoadProgressCallback = None,
        is_cancelled: LoadCancelCheck = None,
    ) -> ExcelSheetMatrix:
        """Capture the sheet as a raw matrix in ONE pass over the file.

        This is the Stage-A replacement for the historical double-read in
        ``import_data`` (an openpyxl pass for fonts plus a pandas pass for
        data). Values keep the raw cell types; legacy font names are stored
        sparsely so the ``auto`` Hindi mode behaves identically.

        Raises:
            DataSourceError: When the file is missing, unreadable, or of an
                unsupported format.
            SheetNotFoundError: When ``sheet_name`` is given but not present.
        """
        if not file_path.exists():
            raise DataSourceError(f"File does not exist: {file_path}")

        suffix = file_path.suffix.lower()

        def cancel_point() -> None:
            if _cancelled(is_cancelled):
                raise LoadCancelled()

        if suffix == ".xlsx":
            matrix = ExcelImporter._load_matrix_xlsx(
                file_path,
                sheet_name,
                header_row,
                hindi_font_mode,
                progress_callback,
                cancel_point,
            )
        elif suffix == ".csv":
            matrix = ExcelImporter._load_matrix_csv(file_path, header_row, progress_callback, cancel_point)
        elif suffix == ".xls":
            matrix = ExcelImporter._load_matrix_pandas_any(
                file_path, sheet_name, header_row, "xlrd", progress_callback, cancel_point
            )
        else:
            raise DataSourceError(f"Unsupported file format: {suffix}")

        matrix.pin_col = ExcelImporter._detect_pin_column(matrix.columns, pin_column_hint)
        return matrix

    @staticmethod
    def _header_cell_name(
        raw_value: Any, font_name: Optional[str], hindi_font_mode: str, c_idx: int
    ) -> str:
        """Convert one header cell into a column name (pandas 'Unnamed' parity)."""
        c_str = str(raw_value).strip() if raw_value is not None else ""
        if raw_value is None or c_str == "" or c_str.startswith("Unnamed:"):
            return f"Unnamed: {c_idx}"
        return HindiConverter.convert_value(c_str, font_name=font_name, mode=hindi_font_mode).strip()

    @staticmethod
    def _cell_font_name(cell: Any) -> Optional[str]:
        """Return the cell's font name when it is a legacy Devanagari font (sparse)."""
        try:
            font = cell.font
            if font is None or not font.name:
                return None
            if HindiConverter.is_krutidev_font(font.name):
                return font.name
        except (AttributeError, TypeError):  # EmptyCell / style lookups
            return None
        return None

    @staticmethod
    def _load_matrix_xlsx(
        file_path: Path,
        sheet_name: Optional[str],
        header_row: int,
        hindi_font_mode: str,
        progress_callback: LoadProgressCallback,
        cancel_point: Callable[[], None],
    ) -> ExcelSheetMatrix:
        """Single read-only openpyxl pass collecting values and legacy fonts."""
        import openpyxl

        matrix = ExcelSheetMatrix(row_index_offset=header_row + 1)
        wb = None
        try:
            wb = openpyxl.load_workbook(file_path, read_only=True, data_only=True)
            if sheet_name:
                if sheet_name not in wb.sheetnames:
                    raise SheetNotFoundError(
                        f"Sheet '{sheet_name}' not found in {file_path.name}. "
                        f"Available: {', '.join(wb.sheetnames)}"
                    )
                ws = wb[sheet_name]
            else:
                ws = wb.active

            max_row = ws.max_row or 0
            total_estimate = max(max_row - header_row - 1, 0)

            max_width = len(matrix.columns)
            for r_idx, row_cells in enumerate(ws.iter_rows()):
                cancel_point()
                if r_idx < header_row:
                    continue
                if r_idx == header_row:
                    header_fonts: Dict[int, str] = {}
                    raw_cells: List[Any] = []
                    for c_idx, cell in enumerate(row_cells):
                        raw_cells.append(cell.value)
                        legacy = ExcelImporter._cell_font_name(cell)
                        if legacy:
                            header_fonts[c_idx] = legacy
                    matrix.columns = [
                        ExcelImporter._header_cell_name(
                            raw, header_fonts.get(c_idx), hindi_font_mode, c_idx
                        )
                        for c_idx, raw in enumerate(raw_cells)
                    ]
                    max_width = max(max_width, len(matrix.columns))
                    continue

                data_row: List[Any] = []
                row_legacy: Dict[int, str] = {}
                for c_idx, cell in enumerate(row_cells):
                    data_row.append(cell.value)
                    if hindi_font_mode != "none":
                        legacy = ExcelImporter._cell_font_name(cell)
                        if legacy:
                            row_legacy[c_idx] = legacy
                max_width = max(max_width, len(data_row))
                data_r = len(matrix.rows)
                if row_legacy:
                    matrix.row_fonts[data_r] = row_legacy
                matrix.rows.append(data_row)
                if len(matrix.rows) % 500 == 0:
                    _progress_shim(progress_callback, len(matrix.rows), total_estimate, "Reading workbook")

            # Pad rows/columns to the widest row so record width matches pandas.
            if matrix.columns and len(matrix.columns) < max_width:
                matrix.columns += [
                    ExcelImporter._header_cell_name(None, None, hindi_font_mode, c_idx)
                    for c_idx in range(len(matrix.columns), max_width)
                ]
            for data_row in matrix.rows:
                if len(data_row) < max_width:
                    data_row.extend([None] * (max_width - len(data_row)))
        except LoadCancelled:
            raise
        except SheetNotFoundError:
            raise
        except DataSourceError:
            raise
        except Exception as e:
            logger.error("Failed to read workbook %s: %s", file_path, e)
            raise DataSourceError(f"Failed to read file: {e}") from e
        finally:
            if wb is not None:
                try:
                    wb.close()
                except Exception:  # pragma: no cover - defensive
                    logger.debug("Workbook close failed", exc_info=True)
        return matrix

    @staticmethod
    def _load_matrix_csv(
        file_path: Path,
        header_row: int,
        progress_callback: LoadProgressCallback,
        cancel_point: Callable[[], None],
    ) -> ExcelSheetMatrix:
        """CSV path: one vectorized pandas read (C parser, already fast)."""
        try:
            df = pd.read_csv(file_path, header=header_row, dtype=object)
        except Exception as e:
            logger.error("Failed to read CSV %s: %s", file_path, e)
            raise DataSourceError(f"Failed to read file: {e}") from e

        columns: List[str] = []
        for c_idx, raw_col in enumerate(df.columns):
            c_str = str(raw_col).strip()
            if c_str.startswith("Unnamed:"):
                columns.append(c_str)
            else:
                columns.append(
                    HindiConverter.convert_value(c_str, font_name=None, mode="auto").strip()
                )

        cancel_point()
        object_rows = df.astype(object).where(df.notna(), "")
        matrix = ExcelSheetMatrix(
            columns=columns,
            rows=object_rows.values.tolist(),
            row_index_offset=header_row + 1,
        )
        _progress_shim(progress_callback, matrix.total_rows, matrix.total_rows, "Reading CSV")
        return matrix

    @staticmethod
    def _load_matrix_pandas_any(
        file_path: Path,
        sheet_name: Optional[str],
        header_row: int,
        engine: str,
        progress_callback: LoadProgressCallback,
        cancel_point: Callable[[], None],
    ) -> ExcelSheetMatrix:
        """Fallback whole-file read for engines without cell font metadata (.xls)."""
        try:
            sheet = sheet_name or 0
            with pd.ExcelFile(file_path, engine=engine) as excel_file:
                if isinstance(sheet, str) and sheet not in excel_file.sheet_names:
                    raise SheetNotFoundError(
                        f"Sheet '{sheet}' not found in {file_path.name}. "
                        f"Available: {', '.join(excel_file.sheet_names)}"
                    )
                df = pd.read_excel(excel_file, sheet_name=sheet, header=header_row, dtype=object)
        except SheetNotFoundError:
            raise
        except Exception as e:
            logger.error("Failed to read spreadsheet %s: %s", file_path, e)
            raise DataSourceError(f"Failed to read file: {e}") from e

        cancel_point()
        columns: List[str] = []
        for c_idx, raw_col in enumerate(df.columns):
            c_str = str(raw_col).strip()
            if c_str.startswith("Unnamed:"):
                columns.append(c_str)
            else:
                columns.append(
                    HindiConverter.convert_value(c_str, font_name=None, mode="auto").strip()
                )
        object_rows = df.astype(object).where(df.notna(), "")
        matrix = ExcelSheetMatrix(
            columns=columns,
            rows=object_rows.values.tolist(),
            row_index_offset=header_row + 1,
        )
        _progress_shim(progress_callback, matrix.total_rows, matrix.total_rows, "Reading spreadsheet")
        return matrix

    # ------------------------------------------------------------------
    # Stage B: chunked row conversion (worker-thread friendly)
    # ------------------------------------------------------------------
    @staticmethod
    def build_chunk_records(
        row_values: List[List[Any]],
        columns: List[str],
        pin_col: Optional[str],
        first_record_row: int,
        hindi_font_mode: str = "auto",
        row_fonts: Optional[Dict[int, Dict[int, str]]] = None,
        start: int = 0,
    ) -> List[SourceRecord]:
        """Convert raw matrix rows ``[start, start+len(row_values))`` into records.

        Replicates the per-cell conversion of ``import_data``
        (NaN -> '', Hindi conversion by legacy font name or auto-detect) with
        ``SourceRecord.model_construct`` so pydantic field validation is
        skipped in worker threads.

        Args:
            row_values: Raw rows of the chunk (list of per-cell values).
            columns: Converted column names of the sheet.
            pin_col: Detected PIN column name.
            first_record_row: ``row_index`` of the anchor row (``header_row + 1``).
            hindi_font_mode: One of ``auto`` / ``kruti_to_unicode`` /
                ``chanakya_to_unicode`` / ``none``.
            row_fonts: Sparse ``{local_row_index: {col_index: legacy_font_name}}``.
            start: Chunk start index within the sheet (0-based data rows).

        Returns:
            The ``SourceRecord`` list of this chunk.
        """
        records: List[SourceRecord] = []
        fonts = row_fonts or {}
        for local_i, row_vals in enumerate(row_values):
            row_dict: Dict[str, Any] = {}
            row_font_map = fonts.get(local_i, {})
            for c_idx, col in enumerate(columns):
                val = row_vals[c_idx] if c_idx < len(row_vals) else None
                if val is None or (isinstance(val, float) and val != val):
                    row_dict[col] = ""
                else:
                    row_dict[col] = HindiConverter.convert_value(
                        val,
                        font_name=row_font_map.get(c_idx),
                        mode=hindi_font_mode,
                    )

            raw_pin_val = row_dict.get(pin_col, "") if pin_col else ""
            raw_pin, normalized_pin = DataNormalizer.normalize_pin(raw_pin_val)

            rec = SourceRecord.model_construct(
                row_index=first_record_row + start + local_i,
                raw_pin=raw_pin,
                pin=normalized_pin,
                data=row_dict,
            )
            records.append(rec)
        return records

    @staticmethod
    def slice_chunk_jobs(
        matrix: ExcelSheetMatrix,
        rows_per_chunk: int,
        hindi_font_mode: str = "auto",
    ) -> List[Dict[str, Any]]:
        """Split a captured matrix into picklable chunk jobs (thread & process safe)."""
        jobs: List[Dict[str, Any]] = []
        total = matrix.total_rows
        row_offset = matrix.row_index_offset
        for start in range(0, total, rows_per_chunk):
            end = min(start + rows_per_chunk, total)
            local_fonts: Dict[int, Dict[int, str]] = {}
            for global_r in range(start, end):
                legacy = matrix.row_fonts.get(global_r)
                if legacy:
                    local_fonts[global_r - start] = legacy
            jobs.append(
                {
                    "start": start,
                    "rows": matrix.rows[start:end],
                    "row_fonts": local_fonts,
                    "columns": matrix.columns,
                    "pin_col": matrix.pin_col,
                    "row_index_offset": row_offset,
                    "hindi_mode": hindi_font_mode,
                }
            )
        return jobs

    @staticmethod
    def _detect_pin_column(columns: List[str], hint: Optional[str] = None) -> Optional[str]:
        """Detect which column represents the PIN Number."""
        if hint and hint in columns:
            return hint

        # Common variants in English and Hindi
        candidates = [
            "pin number",
            "pin_number",
            "pin",
            "id",
            "emp id",
            "employee id",
            "roll no",
            "roll number",
            "badge number",
            "sl.no",
            "sl no",
            "s.no",
            "s no",
            "क्रम",
            "क्रम संख्या",
            "क्र0सं0",
            "क्र०सं०",
            "क्र.सं.",
            "आईडी",
            "पिन",
            "रोल नंबर",
        ]

        # Case-insensitive match
        col_lower_map = {c.lower(): c for c in columns}
        for cand in candidates:
            if cand in col_lower_map:
                return col_lower_map[cand]

        # Partial match
        for cand in ["pin", "id", "क्रम"]:
            for col in columns:
                if cand in col.lower():
                    return col

        return columns[0] if columns else None
